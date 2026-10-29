import base64,io,sys,tempfile,unittest
from contextlib import nullcontext
from types import SimpleNamespace
from datetime import date
from pathlib import Path
from unittest.mock import Mock,patch
from zipfile import ZipFile
import xml.etree.ElementTree as ET
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'app'))
import trip_export as export
import business_trips as trips

def record(day='2026-10-01',category='센터',start='13:00',end='17:00'):
    return dict(date=day,category=category,start=start,end=end,location='출장 장소',purpose='현장 협의',transport='대중교통')

@unittest.skipUnless((ROOT/'app/trip-expenses-template.txt').exists(),'Bundled blank template required')
class ExpenseExport(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.path=Path(self.temp.name)/export.export_filename('2026-10')
        self.fixture=patch.object(export,'TEMPLATE',ROOT/'app/trip-expenses-template.txt');self.fixture.start()
    def tearDown(self):self.fixture.stop();self.temp.cleanup()
    def test_export_job_reads_site_history_without_attendance_lookup_or_write(self):
        app=SimpleNamespace(file_lock=lambda *a,**kw:nullcontext())
        driver=Mock()
        with patch('monthly_reconcile._settings',return_value={}),patch('monthly_reconcile._connect',return_value=(driver,Mock())),patch('monthly_reconcile._read_month') as attendance,patch('trip_site.TripPage') as site:
            site.return_value.existing.return_value=[record(),record('2026-10-02','협회')]
            result=trips.run_job(app,dict(operation='export',month='2026-10',name='시험',output_path=str(self.path)))
            self.assertEqual(result['center_count'],1);self.assertEqual(result['excluded_count'],1)
            attendance.assert_not_called();site.return_value.submit.assert_not_called();driver.quit.assert_called_once()
    def test_export_uses_builtin_even_with_legacy_template_setting(self):
        app=SimpleNamespace(file_lock=lambda *a,**kw:nullcontext());driver=Mock()
        with patch('monthly_reconcile._settings',return_value={}),patch('monthly_reconcile._connect',return_value=(driver,Mock())),patch('trip_site.TripPage') as site:
            site.return_value.existing.return_value=[record()]
            result=trips.run_job(app,dict(operation='export',month='2026-10',name='시험',output_path=str(self.path),template_path='missing-old-local-file.xlsx'))
            self.assertTrue(result['ok']);self.assertEqual(result['center_count'],1)
    def test_export_missing_name_does_not_connect(self):
        with patch('monthly_reconcile._connect') as connect:
            with self.assertRaises(ValueError):trips.run_job(Mock(),dict(operation='export',month='2026-10',name='',output_path=str(self.path)))
            connect.assert_not_called()
    def test_only_center_and_actual_month_and_filename_and_totals(self):
        rows=[record(),record('2026-10-02',start='14:00',end='16:00'),record('2026-10-03','협회'),record('2026-09-01')]
        result=export.export_records(rows,'2026-10',self.path,'시험 사용자')
        self.assertEqual(self.path.name,'2026 출장비지급내역(10월)_.xlsx')
        self.assertEqual((result['center_count'],result['excluded_count'],result['total_amount']),(2,1,30000))
        w=load_workbook(self.path,data_only=True);s=w.active
        self.assertEqual(s['D9'].value,'시험 사용자');self.assertEqual(s['E9'].value,'2026.10.01.~ 2026.10.31.')
        self.assertEqual(s['A12'].value.date(),date(2026,10,1));self.assertEqual(s['A13'].value.date(),date(2026,10,2))
        self.assertIsNone(s['A14'].value)
        self.assertEqual((s['D33'].value,s['D34'].value,s['B35'].value,s['J35'].value),(1,1,30000,30000))
        self.assertEqual(s['C12'].value,1300);self.assertEqual(s['D12'].value,1700)
        self.assertEqual(load_workbook(self.path,data_only=False).active['J35'].value,'=B35')
    def test_native_template_styles_merges_dimensions_and_printer_preserved(self):
        export.export_records([record()],'2026-10',self.path,'이름')
        with ZipFile(io.BytesIO(base64.b64decode(export.TEMPLATE.read_bytes()))) as template,ZipFile(self.path) as result:
            for part in ['xl/styles.xml','xl/theme/theme1.xml','xl/printerSettings/printerSettings1.bin']:
                if part in template.namelist():self.assertEqual(template.read(part),result.read(part))
            a=ET.fromstring(template.read('xl/worksheets/sheet1.xml'));b=ET.fromstring(result.read('xl/worksheets/sheet1.xml'))
            for name in ['cols','mergeCells','pageMargins','pageSetup','printOptions']:
                aa=a.find('s:'+name,export.NS);bb=b.find('s:'+name,export.NS)
                self.assertEqual(ET.tostring(aa) if aa is not None else None,ET.tostring(bb) if bb is not None else None,name)
            for aa,bb in zip(a.findall('.//s:row',export.NS),b.findall('.//s:row',export.NS)):
                self.assertEqual(aa.attrib,bb.attrib)
    def test_nineteen_row_capacity_continues_on_identical_form(self):
        rows=[record(f'2026-10-{n:02}') for n in range(1,23)]
        result=export.export_records(rows,'2026-10',self.path,'시험')
        w=load_workbook(self.path,data_only=True)
        self.assertEqual(result['sheet_count'],2);self.assertEqual(len(w.worksheets),2)
        self.assertEqual(w.worksheets[1]['A12'].value.date(),date(2026,10,20))
        self.assertEqual([s['B35'].value for s in w],[380000,60000])
        self.assertTrue(all(str(s.print_area).endswith("!$A$1:$L$35") for s in w))
    def test_daily_event_duration_uses_actual_hours_and_capped_amount(self):
        export.export_records([record(start='09:00',end='18:00')],'2026-10',self.path,'이름')
        s=load_workbook(self.path,data_only=True).active
        self.assertEqual(s['F12'].value.total_seconds(),9*3600)
        self.assertEqual(s['B35'].value,20000)
    def test_empty_unknown_category_duplicate_and_over_cap_do_not_export(self):
        cases=[[record(category='협회')],[record(category='협의')],[record(),record()],
               [record(start='09:00',end='13:00'),record(start='14:00',end='18:00')]]
        for rows in cases:
            with self.assertRaises(ValueError):export.export_records(rows,'2026-10',self.path,'이름')
            self.assertFalse(self.path.exists())
    def test_literal_purpose_never_becomes_a_formula(self):
        row=record();row['purpose']='=HYPERLINK("bad")'
        export.export_records([row],'2026-10',self.path,'이름')
        c=load_workbook(self.path,data_only=False).active['K12']
        self.assertEqual(c.data_type,'s');self.assertEqual(c.value,row['purpose'])
    def test_built_in_template_has_no_device_or_source_metadata(self):
        with ZipFile(io.BytesIO(base64.b64decode(export.TEMPLATE.read_bytes()))) as z:
            self.assertFalse(any(n.startswith('xl/printerSettings/') for n in z.namelist()))
            self.assertEqual(len(ET.fromstring(z.read('docProps/core.xml'))),0)
            self.assertNotIn('absPath',z.read('xl/workbook.xml').decode('utf-8'))
    def test_original_template_cannot_be_overwritten(self):
        original=Path(self.temp.name)/'original.xlsx';before=base64.b64decode(export.TEMPLATE.read_bytes());original.write_bytes(before)
        with self.assertRaisesRegex(ValueError,'원본 양식'):
            export.export_records([record()],'2026-10',original,'이름',template=original)
        self.assertEqual(original.read_bytes(),before)
    def test_template_contains_no_original_employee_or_historical_records(self):
        with ZipFile(io.BytesIO(base64.b64decode(export.TEMPLATE.read_bytes()))) as z:
            root=ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
            name=root.find('.//s:c[@r="D9"]',export.NS)
            self.assertFalse(''.join(name.itertext()).strip())
            self.assertTrue(root.findall('.//s:c[@r="A1"]/s:is/s:r/s:rPr',export.NS))
            self.assertNotIn('xl/sharedStrings.xml',z.namelist())

class CategoryPlanning(unittest.TestCase):
    def test_center_default_and_existing_association_is_preserved(self):
        rows=trips.make_rows([dict(date='2026-10-01',checkin='08:40',checkout='18:40')],{},[],date(2026,10,29))
        self.assertEqual(rows[0]['category'],'센터')
        changed=trips.merge_entries(rows,[dict(rows[0],selected=True,category='협회')])[0]
        self.assertEqual(changed['category'],'협회')
        with self.assertRaises(ValueError):trips.merge_entries(rows,[dict(rows[0],selected=True,category='기타')])
        existing=trips.make_rows([dict(date='2026-10-01',checkin='08:40',checkout='18:40')],{},[record(category='협회')],date(2026,10,29))
        self.assertEqual(existing[0]['category'],'협회')

if __name__=='__main__':unittest.main()
