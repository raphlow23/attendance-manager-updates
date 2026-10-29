"""Real Chrome workflows against a local trip/overtime form+POST server."""
import sys,threading,tempfile,re
from pathlib import Path
from datetime import date
from html import escape
from http.server import BaseHTTPRequestHandler,ThreadingHTTPServer
from urllib.parse import urlsplit,parse_qs
from unittest.mock import patch,Mock
ROOT=Path(__file__).resolve().parents[1];sys.path[:0]=[str(ROOT/'.test-libs'),str(ROOT/'app')]
from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.remote.webelement import WebElement
import main as app
import business_trips as trips
import overtime_requests as overtime

trip_records=[];ot_records=[];posts=[];ignore_ot_update=False;direct_trip_landing=False
NAV='<a href="/attendance">근태관리</a><a href="/work">근무일지</a>'
def inp(name,value='',placeholder=''):return f'<input name="{name}" value="{escape(str(value),quote=True)}" placeholder="{placeholder}">'
def sel(name,options):return '<select name="'+name+'">'+''.join(f'<option value="{v}">{label}</option>' for v,label in options)+'</select>'
def numbers(name,limit,unit):return sel(name,[(str(n),str(n)+unit) for n in range(limit)])
def trip_page(edit=None):
    saved=trip_records[edit] if edit is not None else None
    form='<form method="post" action="/trip-save">'+inp('nonce','fixture-token') .replace('<input ','<input type="hidden" ',1)
    form+=sel('kind',[('x','선택'),('center','센터'),('association','협회')])+inp('y',2026)+inp('m',10)+inp('d',27)+inp('place','','출장장소')+inp('purpose','','출장목적을 구체적으로 입력하세요')
    form+=sel('transport',[('x','선택'),('public','대중교통')])+numbers('sh',24,'시')+numbers('sm',60,'분')+numbers('eh',24,'시')+numbers('em',60,'분')+'<button type="submit">확인</button></form>'
    heads=''.join('<th>'+t+'</th>' for t in ('구분','출장일','출장장소','출장목적','이동수단','출발시간','복귀시간','출장시간'))
    body=''.join('<tr>'+''.join('<td>'+escape(str(v))+'</td>' for v in (r.get('category','센터'),f"{int(r['date'][5:7])}월 {int(r['date'][8:])}일 (수)",r['location'],r['purpose'],'대중교통',r['start'],r['end'],'시간'))+'</tr>' for r in trip_records)
    if saved:
        form=form.replace('action="/trip-save"','action="/trip-update?index='+str(edit)+'"')
        for name,value in dict(y=int(saved['date'][:4]),m=int(saved['date'][5:7]),d=int(saved['date'][8:]),place=saved['location'],purpose=saved['purpose']).items():
            placeholder='출장장소' if name=='place' else '출장목적을 구체적으로 입력하세요' if name=='purpose' else ''
            previous=2026 if name=='y' else 10 if name=='m' else 27 if name=='d' else ''
            form=form.replace(inp(name,previous,placeholder),inp(name,value,placeholder))
        for name,value in dict(kind='association' if saved.get('category')=='협회' else 'center',transport='public',sh=int(saved['start'][:2]),sm=int(saved['start'][3:]),eh=int(saved['end'][:2]),em=int(saved['end'][3:])).items():
            a=form.index('<select name="'+name+'">');b=form.index('</select>',a)
            form=form[:a]+form[a:b].replace('value="'+str(value)+'"','value="'+str(value)+'" selected')+form[b:]
        # Reproduce saved option + generated list duplicates on the real site.
        for name,value in dict(sh=int(saved['start'][:2]),sm=int(saved['start'][3:]),eh=int(saved['end'][:2]),em=int(saved['end'][3:])).items():
            form=form.replace('<select name="'+name+'">','<select name="'+name+'"><option value="'+str(value)+'" selected>'+str(value)+'</option>',1)
        form=form.replace('>확인</button>','>수정</button>')
        # Actual edit screen: column headers over inputs, no placeholders.
        def field(name):
            pattern=r'<select name="'+name+r'">.*?</select>' if name in ('kind','transport','sh','sm','eh','em') else r'<input name="'+name+r'"[^>]*>'
            return re.search(pattern,form).group().replace('placeholder="출장장소"','').replace('placeholder="출장목적을 구체적으로 입력하세요"','')
        opening=form[:form.index('>')+1]
        hidden=re.search(r'<input type="hidden"[^>]*>',form).group()
        groups=[('구분',('kind',)),('출장일',('y','m','d')),('출장장소',('place',)),('출장목적',('purpose',)),('이동수단',('transport',)),('출발시간',('sh','sm')),('복귀시간',('eh','em'))]
        values=[''.join(field(n) for n in names) for _,names in groups]
        form=opening+hidden+'<table><thead><tr>'+''.join('<th>'+h+'</th>' for h,_ in groups)+'<th>수정</th></tr></thead><tbody><tr>'+''.join('<td>'+v+'</td>' for v in values)+'<td><button type="submit">수정</button></td></tr></tbody></table></form>'
    body=body.replace('</tr>','<td>EDIT</td></tr>')
    for index in range(len(trip_records)):body=body.replace('EDIT','<a href="/trip-edit?index='+str(index)+'">수정</a> <a href="/trip-delete?index='+str(index)+'" onclick="if(!confirm(\'삭제할까요?\'))return false;alert(\'삭제를 진행합니다\');return true">삭제</a>',1)
    return NAV+'<h2>출장 등록하기</h2>'+((trip_page().split('<h2>출장 등록하기</h2>')[1].split('<h2>사용자님')[0]) if saved else '')+form+'<h2>사용자님의 10월 출장일지</h2><table><thead><tr>'+heads+'</tr></thead><tbody>'+body+'</tbody></table>'
def overtime_page(edit=None):
    r=ot_records[edit] if edit is not None else None
    form='<form method="post" action="'+('/ot-update?index='+str(edit) if r else '/ot-save')+'">'
    form+=(inp('dept','센터') if r else sel('dept',[('x','선택'),('center','센터')]))+inp('purpose',r['purpose'] if r else '','근무목적을 구체적으로 입력하세요')+inp('year',2026)
    form+=numbers('month',13,'월')+numbers('day',32,'일')+sel('weekday',[(v,v+'요일') for v in '월화수목금토일'])
    form+=numbers('sh',24,'시')+numbers('sm',60,'분')+numbers('eh',24,'시')+numbers('em',60,'분')
    form+=inp('fy',2026)+inp('fm',int(r['date'][5:7]) if r else 10)+inp('fd',int(r['date'][8:]) if r else 27)+'<button type="submit">'+('시간외근무 수정하기' if r else '시간외근무 신청')+'</button></form>'
    if r:
        for name,value in dict(month=int(r['date'][5:7]),day=int(r['date'][8:]),sh=int(r['start'][:2]),sm=int(r['start'][3:]),eh=int(r['end'][:2]),em=int(r['end'][3:])).items():
            marker='<select name="'+name+'">';a=form.index(marker);b=form.index('</select>',a)
            form=form[:a]+form[a:b].replace('value="'+str(value)+'"','value="'+str(value)+'" selected')+form[b:]
    if r:form=form.replace(inp('purpose',r['purpose'],'근무목적을 구체적으로 입력하세요'),'<table><tr><td>근무목적 :</td><td>'+inp('purpose',r['purpose'])+'</td></tr></table>')
    body=''.join('<tr><td>'+escape(r['date']+' (목)')+'</td><td>'+escape(r['start']+' ~ '+r['end']+' (1시간 00분)')+'</td><td><a href="/ot-edit?index='+str(i)+'">수정하기</a></td></tr>' for i,r in enumerate(ot_records))
    return NAV+'<h2>시간외근무 '+('수정' if r else '신청서')+'</h2>'+form+'<table><thead><tr><th>근무일</th><th>근무시간</th><th>수정</th></tr></thead><tbody>'+body+'</tbody></table>'
class Handler(BaseHTTPRequestHandler):
    def log_message(self,*args):pass
    def do_GET(self):
        path=urlsplit(self.path).path
        if path=='/trip-delete':
            removed=trip_records.pop(int(parse_qs(urlsplit(self.path).query)['index'][0]));posts.append(('trip-delete',removed['date']))
            self.send_response(303);self.send_header('Location','/trips');self.end_headers();return
        if path=='/work' and direct_trip_landing:content=trip_page()
        elif path=='/work':content=NAV+'<div><span><h3>출장일지</h3></span><a href="/trips">바로가기</a></div><div><span><h3>시간외근무신청</h3></span><a href="/overtime">바로가기</a></div>'
        elif path in ('/trips','/sappd_g6.php'):content=trip_page()
        elif path=='/trip-edit':content=trip_page(int(parse_qs(urlsplit(self.path).query)['index'][0]))
        elif path=='/overtime':content=overtime_page()
        elif path=='/ot-edit':content=overtime_page(int(parse_qs(urlsplit(self.path).query)['index'][0]))
        elif path=='/ot-success':content=NAV+'<h2>신청 저장 완료</h2>'
        elif path=='/attendance':
            month=int(parse_qs(urlsplit(self.path).query).get('m',['10'])[0])
            options=''.join(f'<option value="{n}"'+(' selected' if n==month else '')+f'>{n}월</option>' for n in range(1,13))
            content=NAV+f'<h2>{month}월 근태현황</h2><form method="get" action="/attendance"><select name="m">'+options+'</select><button type="submit">조회</button></form>'
        else:self.send_response(404);self.end_headers();return
        self.send_response(200);self.send_header('Content-Type','text/html; charset=utf-8');self.end_headers();self.wfile.write(('<meta charset="utf-8">'+content).encode())
    def do_POST(self):
        fields={k:v[0] for k,v in parse_qs(self.rfile.read(int(self.headers['Content-Length'])).decode(),keep_blank_values=True).items()}
        if urlsplit(self.path).path in ('/trip-save','/trip-update'):
            assert fields['kind'] in ('center','association') and fields['transport']=='public' and fields['nonce']=='fixture-token'
            row=dict(category='센터' if fields['kind']=='center' else '협회',date=f"{int(fields['y']):04}-{int(fields['m']):02}-{int(fields['d']):02}",location=fields['place'],purpose=fields['purpose'],start=f"{int(fields['sh']):02}:{int(fields['sm']):02}",end=f"{int(fields['eh']):02}:{int(fields['em']):02}")
            if urlsplit(self.path).path=='/trip-update':
                trip_records[int(parse_qs(urlsplit(self.path).query)['index'][0])]=row;posts.append(('trip-update',row['date']))
            else:trip_records.append(row);posts.append(('trip',row['date']))
            target='/trips'
        elif urlsplit(self.path).path in ('/ot-save','/ot-update'):
            updating=urlsplit(self.path).path=='/ot-update'
            assert fields['dept']==('센터' if updating else 'center') and fields['year']==fields['fy'] and fields['month']==fields['fm'] and fields['day']==fields['fd']
            row=dict(date=f"{int(fields['year']):04}-{int(fields['month']):02}-{int(fields['day']):02}",purpose=fields['purpose'],start=f"{int(fields['sh']):02}:{int(fields['sm']):02}",end=f"{int(fields['eh']):02}:{int(fields['em']):02}")
            if updating:
                if not ignore_ot_update:ot_records[int(parse_qs(urlsplit(self.path).query)['index'][0])]=row
            else:ot_records.append(row)
            posts.append(('overtime-update' if updating else 'overtime',row['date']));target='/ot-success'
        else:self.send_response(404);self.end_headers();return
        self.send_response(303);self.send_header('Location',target);self.end_headers()
class Server(ThreadingHTTPServer):
    def handle_error(self,request,address):
        if not isinstance(sys.exc_info()[1],(ConnectionResetError,BrokenPipeError)):super().handle_error(request,address)
server=Server(('127.0.0.1',0),Handler);thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start();base=f'http://127.0.0.1:{server.server_port}'
def connect(*args):
    options=webdriver.ChromeOptions()
    options.set_capability("unhandledPromptBehavior","ignore")
    for arg in ('--headless=new','--window-size=1440,1000','--disable-background-networking','--no-first-run'):options.add_argument(arg)
    driver=webdriver.Chrome(options=options,service=Service(str(ROOT/'app/driver/chromedriver.exe'),popen_kw={'creation_flags':app.subprocess.CREATE_NO_WINDOW}))
    driver.get(base+'/attendance');return driver,Mock()
try:
 with tempfile.TemporaryDirectory(prefix='workday-journal-') as folder:
  with patch.object(app,'DATA_DIR',folder),patch.object(app,'STATE_FILE',str(Path(folder)/'state.json')),patch.object(app,'SETTINGS_FILE',str(Path(folder)/'settings.json')),patch.object(app,'LOG_FILE',str(Path(folder)/'events.log')):
   app.save_json(app.SETTINGS_FILE,{'id':'fixture','pw':'fixture'});app.save_json(app.STATE_FILE,{'day_rules':{'2026-10-07':{'exclude_overtime':True}}})
   attendance=[dict(date='2026-10-07',checkin='09:00',checkout='18:00',note=''),dict(date='2026-10-13',checkin='08:35',checkout='18:00',note='')]
   original=WebElement.get_attribute
   def initial_value(element,name):return element.get_dom_attribute(name) if name=='value' else original(element,name)
   with patch.object(WebElement,'click',lambda self:None),patch.object(WebElement,'send_keys',lambda self,*args:None),patch.object(WebElement,'get_attribute',initial_value),patch('monthly_reconcile._connect',side_effect=connect),patch('monthly_reconcile._read_month',return_value=attendance):
    review=trips.run_job(app,dict(operation='scan',month='2026-10',total=30000))
    chosen=[]
    for row in review['rows']:chosen.append(dict(date=row['date'],selected=True,hours=4 if row.get('event_day') else 2,start_hour=9,location='장소 '+row['date'],purpose='직접  입력한 출장 목적'))
    review=trips.run_job(app,dict(operation='inspect',review_id=review['review_id'],total=30000,rows=chosen))
    assert review['ready']
    result=trips.run_job(app,dict(operation='submit',review_id=review['review_id'],total=30000,rows=chosen))
    assert result['ok'] and result['summary']['remaining_amount']==0
    assert posts==[('trip','2026-10-07'),('trip','2026-10-13')]
    assert trip_records[0]['start']=='09:00' and trip_records[0]['end']=='18:00' and trip_records[1]['end']=='11:00'
    try:trips.run_job(app,dict(operation='submit',review_id=review['review_id'],total=30000,rows=chosen));raise AssertionError('Review replay allowed')
    except ValueError:pass
    for key in ('2026-10-14','2026-10-15','2026-10-16'):attendance.append(dict(date=key,checkin='08:35',checkout='18:00',note=''))
    under=trips.run_job(app,dict(operation='scan',month='2026-10',total=80000))
    under_choices=[dict(date=r['date'],selected=r['date']=='2026-10-14',hours=2 if r['date']=='2026-10-14' else r['hours'],start_hour=r['start_hour'],location='부족 일정 장소',purpose='부족 일정 목적') for r in under['rows']]
    before_count=len(posts)
    under=trips.run_job(app,dict(operation='inspect',review_id=under['review_id'],total=80000,rows=under_choices))
    assert under['ready'] and under['summary']['remaining_amount']==40000 and len(posts)==before_count,under
    partial=trips.run_job(app,dict(operation='scan',month='2026-10',total=80000))
    plans=[dict(date=r['date'],selected=r['date']>='2026-10-14',hours=2 if r['date']=='2026-10-14' else 4,start_hour=13,location='직접 입력 장소',purpose='등록할 목적') for r in partial['rows']]
    partial=trips.run_job(app,dict(operation='inspect',review_id=partial['review_id'],total=80000,rows=plans))
    assert partial['ready'],partial
    for r in plans:
     if r['date']=='2026-10-15':r['purpose']='   '
     if r['date']=='2026-10-16':r['selected']=False
    result=trips.run_job(app,dict(operation='submit',review_id=partial['review_id'],total=80000,rows=plans))
    assert result['ok'] and [r['date'] for r in result['rows']]==['2026-10-14'] and result['summary']['remaining_amount']==40000,result
    assert [r['date'] for r in trip_records]==['2026-10-07','2026-10-13','2026-10-14']
    loaded=trips.run_job(app,dict(operation='scan',month='2026-10',total=40000))
    loaded=trips.run_job(app,dict(operation='inspect',review_id=loaded['review_id'],total=40000,rows=[],edit_existing=True))
    assert loaded['ready'] and sum(r['selected'] for r in loaded['rows'])==3,loaded
    edits=[dict(r,selected=r['date']=='2026-10-14') for r in loaded['rows']]
    edit=next(r for r in edits if r['selected']);assert edit['location']=='직접 입력 장소' and edit['purpose']=='등록할 목적'
    edit.update(location='수정한 장소',purpose='수정한 출장 목적',start_hour=12,start_minute=30,start='12:30',end='14:30',category='협회')
    edited=trips.run_job(app,dict(operation='submit',review_id=loaded['review_id'],total=40000,rows=edits))
    assert edited['ok'] and edited['rows'][0]['status']=='수정 확인 완료',edited
    assert trip_records[-1]['category']=='협회'
    assert len(trip_records)==3 and trip_records[-1]['location']=='수정한 장소' and trip_records[-1]['purpose']=='수정한 출장 목적'
    assert (trip_records[-1]['start'],trip_records[-1]['end'])==('12:30','14:30')
    assert [r for r in posts if r[0]=='trip-update']==[('trip-update','2026-10-14')]
    export_path=Path(folder)/'2026 출장비지급내역(10월)_.xlsx'
    exported=trips.run_job(app,dict(operation='export',month='2026-10',template_path=str(ROOT/'app/trip-expenses-template.txt'),output_path=str(export_path),name='시험 사용자',position='대리'))
    assert exported['ok'] and exported['center_count']==2 and exported['excluded_count']==1 and exported['total_amount']==30000,exported
    from openpyxl import load_workbook
    report=load_workbook(export_path,data_only=True).active
    assert report['B35'].value==30000 and report['D9'].value=='시험 사용자'
    batch_scan=trips.run_job(app,dict(operation='scan',month='2026-10',total=50000))
    batch_rows=[dict(r,pending_delete=r['date']=='2026-10-13',selected=r['date']=='2026-10-15',location='일괄 등록 장소' if r['date']=='2026-10-15' else r['location'],purpose='일괄 등록 목적' if r['date']=='2026-10-15' else r['purpose']) for r in batch_scan['rows']]
    before_posts=list(posts)
    batch_review=trips.run_job(app,dict(operation='inspect',review_id=batch_scan['review_id'],total=50000,rows=batch_rows))
    assert batch_review['ready'] and posts==before_posts and any(r['date']=='2026-10-13' for r in trip_records),batch_review
    batch_done=trips.run_job(app,dict(operation='submit',review_id=batch_review['review_id'],total=50000,rows=batch_review['rows']))
    assert batch_done['ok'],batch_done
    assert not any(r['date']=='2026-10-13' for r in trip_records) and any(r['date']=='2026-10-15' for r in trip_records)
    assert [p for p in posts[len(before_posts):] if p[0] in ('trip-delete','trip')]==[('trip-delete','2026-10-13'),('trip','2026-10-15')]
    scan=trips.run_job(app,dict(operation='scan',month='2026-10',total=50000))
    removed=trips.run_job(app,dict(operation='delete',review_id=scan['review_id'],dates=['2026-10-14'],confirmed=True,total=50000,rows=scan['rows']))
    assert removed['ok'] and len(trip_records)==2,removed
    available=next(r for r in removed['rows'] if r['date']=='2026-10-14')
    assert not available['is_existing'] and not available['selected'] and not available['blocked'] and available['existing_amount']==0,available
    again=[dict(r,selected=r['date']=='2026-10-14',purpose='삭제 후 재등록 목적' if r['date']=='2026-10-14' else r['purpose']) for r in removed['rows']]
    inspected=trips.run_job(app,dict(operation='inspect',review_id=removed['review_id'],total=50000,rows=again))
    assert inspected['ready'],inspected
    registered=trips.run_job(app,dict(operation='submit',review_id=inspected['review_id'],total=50000,rows=inspected['rows']))
    assert registered['ok'] and len(trip_records)==3,registered
    assert sum(r['date']=='2026-10-14' for r in trip_records)==1
    assert next(r for r in trip_records if r['date']=='2026-10-14')['purpose']=='삭제 후 재등록 목적'
    assert [r for r in posts if r[0]=='trip-delete']==[('trip-delete','2026-10-13'),('trip-delete','2026-10-14')]
    driver,_=connect()
    try:
     from attendance_recovery import AttendancePage
     attendance_page=AttendancePage(driver,date(2026,10,27));attendance_page.open();attendance_page.choose_month(date(2026,9,1))
     assert attendance_page.month_heading()==(None,9)
     attendance_page.choose_month(date(2026,10,1));assert attendance_page.month_heading()==(None,10)
     direct_trip_landing=True
     from trip_site import TripPage
     TripPage(driver).open()
     assert driver.current_url==base+'/work' and TripPage(driver).form()
     direct_trip_landing=False
     with patch('trip_site.TRIP_ENTRY_URL',base+'/sappd_g6.php'):
      TripPage(driver).open()
      assert driver.current_url==base+'/sappd_g6.php' and TripPage(driver).form()
     page=overtime.OvertimePage(driver);page.open();rows=[]
     for d,checkout in ((1,'19:55'),(7,'20:30')):rows.append(dict(overtime.application_row(date(2026,10,d),{'checkin':'08:35','checkout':checkout}),purpose='직접  입력한 시간외 목적'))
     statuses={};result=overtime.submit_rows(page,rows,rows,statuses,lambda values:None)
     assert all(row['status']=='신청 확인 완료' for row in result),result
     changed=[dict(rows[0],purpose='수정한 새 근무목적')]
     updated=overtime.submit_rows(page,changed,rows,statuses,lambda values:None)
     assert updated[0]['status']=='수정 확인 완료',updated
     assert len(ot_records)==2 and ot_records[0]['end']=='19:00' and ot_records[1]['end']=='20:00'
     assert [p for p in posts if p[0]=='overtime']==[('overtime','2026-10-01'),('overtime','2026-10-07')]
     assert ot_records[0]['purpose']=='수정한 새 근무목적' and ot_records[1]['purpose']=='직접  입력한 시간외 목적'
     assert len([p for p in posts if p[0]=='overtime-update'])==1
     missing=dict(overtime.application_row(date(2026,10,6),{'checkin':'08:35','checkout':'19:15'}),purpose='이전 확인 실패지만 실제 미등록')
     statuses[missing['date']]={'status':'uncertain'}
     added=overtime.submit_rows(page,[missing],rows+[missing],statuses,lambda values:None)
     assert added[0]['status']=='신청 확인 완료' and len(ot_records)==3,added
     assert page.existing(rows[0]) and not page.existing(rows[0],True)
     ignore_ot_update=True
     rejected=overtime.submit_rows(page,[dict(rows[0],purpose='저장되지 않는 목적'),rows[1]],rows,statuses,lambda values:None)
     assert len(rejected)==1 and rejected[0]['status'].startswith('확인 필요:') and statuses[rows[0]['date']]['status']=='uncertain',rejected
     assert len(ot_records)==3 and ot_records[0]['purpose']=='수정한 새 근무목적'
    finally:driver.quit()
   print('JOURNAL_BROWSER_OK: attendance month 10→9→10 with native clicks disabled; underfilled schedule details allowed without POST; trip actual header-row edit form without placeholders matched by column; trip existing category/time/place/purpose updated / no duplicate; center-only XLSX export from live history excludes association; deferred deletion plus registration in one session; two successive site dialogs accepted; trip deletion through observed row button and site confirm -> verified absence -> unregistered row -> one new registration; trip short-date history after POST verified / no duplicate; trip workflow + unchecked/blank-purpose exclusions and partial-budget warning + work-log direct form + user-supplied direct URL without menu links; overtime purpose absent from list / success page without history / observed-page refresh / two sequential POSTs / existing date purpose updated once / uncertain missing date added / detail purpose verified / no duplicate dates; no target-site requests')
finally:server.shutdown();server.server_close();thread.join()
